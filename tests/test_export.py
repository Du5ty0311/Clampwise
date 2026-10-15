"""Exports (B14): the workbook and the tidy CSV read back: every recording-level
column carries its definition as a comment, the numbers are the results'."""

from __future__ import annotations

import csv
import json

import openpyxl
import pytest

from clampwise import export, simulate
from clampwise.analysis import registry
from clampwise.project import Project


@pytest.fixture
def project(tmp_path):
    folder = tmp_path / "A1_61826"
    folder.mkdir()
    p = Project.open(folder)
    sims = {"steps.abf": simulate.steps(), "evoked.abf": simulate.evoked(n_sweeps=10)}
    for name in sims:
        (folder / name).write_bytes(name.encode())          # the store hashes a file; the samples are the simulator's
    p.add_recordings([folder / n for n in sims])
    runs = {n: p.analyze(folder / n, recording=s.recording) for n, s in sims.items()}
    yield p, runs
    p.close()


def _columns(ws) -> dict[str, int]:
    return {c.value: c.column for c in ws[2] if c.value}


def test_the_workbook_has_a_sheet_per_type_with_definitions_as_comments(project, tmp_path):
    p, runs = project
    out = export.write_workbook(p, tmp_path / "out.xlsx")
    book = openpyxl.load_workbook(out)
    assert {"Current steps", "Evoked, voltage clamp", "Definitions", "Provenance"} <= set(book.sheetnames)
    assert any(name.startswith("steps (intrinsic)") for name in book.sheetnames)
    ws = book["Current steps"]
    cols = _columns(ws)
    rin = registry.load()["intrinsic"].output("r_input")
    assert ws.cell(2, cols[rin.title]).comment.text.startswith(rin.definition)
    assert ws.cell(3, cols[rin.title]).value == rin.units
    assert ws.cell(4, cols["File"]).value == "steps.abf"
    assert ws.cell(4, cols[rin.title]).value == pytest.approx(runs["steps.abf"].results["intrinsic"].values["r_input"])
    for sheet in book.worksheets:                             # every column of every sheet is defined
        if sheet.title in ("Definitions", "Provenance"):
            continue
        for c in sheet[2]:
            if c.value:
                assert c.comment is not None and "no definition" not in c.comment.text, (sheet.title, c.value)
    ev = book["Evoked, voltage clamp"]
    amp = registry.load()["evoked"].output("p1_amplitude")
    assert ev.cell(4, _columns(ev)[amp.title]).value == pytest.approx(
        runs["evoked.abf"].results["evoked"].values["p1_amplitude"])
    assert "Choices" in _columns(ev)
    keys = {row[1].value for row in book["Definitions"].iter_rows(min_row=2)}
    assert "intrinsic.r_input" in keys and "evoked.p1_amplitude" in keys


def test_tidy_csv_holds_the_same_numbers_with_a_manifest(project, tmp_path):
    p, runs = project
    written = export.write_csv(p, tmp_path / "csv")
    names = {w.name for w in written}
    assert {"values.csv", "definitions.csv", "flags.csv", "provenance.csv", "manifest.json"} <= names
    with open(tmp_path / "csv" / "values.csv", encoding="utf-8-sig") as f:
        rows = list(csv.DictReader(f))
    r_in = next(r for r in rows if r["File"] == "steps.abf" and r["Analysis"] == "intrinsic" and r["Output"] == "r_input")
    assert float(r_in["Value"]) == pytest.approx(runs["steps.abf"].results["intrinsic"].values["r_input"])
    assert r_in["Units"] == "MOhm"
    manifest = json.loads((tmp_path / "csv" / "manifest.json").read_text(encoding="utf-8"))
    assert manifest["recordings"] == 2 and "values.csv" in manifest["files"]


def test_provenance_carries_each_file_s_hash(project, tmp_path):
    p, runs = project
    export.write_csv(p, tmp_path / "csv")
    with open(tmp_path / "csv" / "provenance.csv", encoding="utf-8-sig") as f:
        rows = {r["File"]: r for r in csv.DictReader(f)}
    for row in p.recordings():
        assert len(row.sha256) == 64 and rows[row.name]["SHA-256"] == row.sha256
