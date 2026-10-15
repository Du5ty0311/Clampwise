"""The figure builder without its window (B12, B23): one value per cell, the
baseline and bins, pulse ratios, the caption and the warnings, comparisons at
the right level, every export reopening editable, and tables from Excel and
from a Clampwise workbook."""

from __future__ import annotations

import numpy as np
import pandas as pd
import pytest

from clampwise.figure import data, prepare, render, stats
from clampwise.figure.recipe import Recipe, Spec, pack


def nested(effect: float = 0.0, animals: int = 4, cells: int = 3, events: int = 20, seed: int = 0) -> pd.DataFrame:
    """Events in cells in animals, control and drug (different animals)."""
    rng = np.random.default_rng(seed)
    rows = []
    for g, shift in (("control", 0.0), ("drug", effect)):
        for a in range(animals):
            am = rng.normal(0, 1)
            for c in range(cells):
                cm = am + rng.normal(0, 1) + shift
                for v in cm + rng.normal(0, 2, events):
                    rows.append({"Condition": g, "Animal": f"{g}-A{a}", "Cell": f"{g}-A{a}-C{c}",
                                 "Amplitude (pA)": -20 + v})
    return pd.DataFrame(rows)


def test_one_value_per_cell_and_a_caption_that_says_so():
    frame = nested()
    r = Recipe(question="compare", plot="points", y="Amplitude (pA)", group="Condition", cell="Cell", animal="Animal")
    p = prepare.prepare(frame, r)
    assert len(p.points) == 24 and set(p.points.rows) == {20}
    assert p.counts["control"] == (12, 4, 240)
    assert "n = 12 cells from 4 animals (240 rows)" in p.caption and "mean ± SEM across cells: how precisely" in p.caption
    assert not p.warnings
    rows = prepare.prepare(frame, Recipe(**{**r.__dict__, "level": "row"}))
    assert len(rows.points) == 480 and "pseudoreplication" in rows.warnings[0]
    small = prepare.prepare(frame[frame.Cell.isin(["control-A0-C0", "control-A0-C1", "drug-A1-C0",
                                                   "drug-A1-C1", "drug-A1-C2"])], r)
    assert any("control has n = 2" in w for w in small.warnings)


def test_baseline_bins_and_pulse_ratios_are_exact():
    t = np.arange(0, 40, 0.5)
    rows = [{"Cell": c, "Time (min)": x, "Slope": (10.0 if x < 10 else 15.0) * (1 + c)} for c in range(3) for x in t]
    r = Recipe(question="time", plot="time_course", x="Time (min)", y="Slope", cell="Cell", normalize=[0, 10],
               bin=2.0)
    p = prepare.prepare(pd.DataFrame(rows), r)
    s = p.summary.set_index("x")
    assert s.loc[1.0, "centre"] == pytest.approx(100.0) and s.loc[21.0, "centre"] == pytest.approx(150.0)
    assert s.loc[21.0, "n"] == 3 and "% of each cell's own mean from 0 to 10" in p.caption
    assert "bins of 2" in p.caption
    pulses = [{"Cell": c, "Sweep": k, "Pulse": j, "Amplitude": -100.0 * (1 + c) * (1, 1.5, 0.8)[j - 1]}
              for c in range(4) for k in range(5) for j in (1, 2, 3)]
    p = prepare.prepare(pd.DataFrame(pulses), Recipe(question="pulses", plot="profile", x="Pulse", y="Amplitude",
                                                     cell="Cell"))
    assert list(p.summary.centre) == pytest.approx([1.0, 1.5, 0.8])
    assert (p.summary.lo == p.summary.centre).all()                 # every cell the same ratio: no spread


def test_filters_keep_only_the_rows_asked_for():
    frame = nested()
    r = Recipe(question="compare", y="Amplitude (pA)", group="Condition", cell="Cell",
               filters=[["Condition", "is", "drug"], ["Amplitude (pA)", "<", "-20"]])
    p = prepare.prepare(frame, r)
    assert set(p.rows.group) == {"drug"} and (p.rows.y < -20).all()
    assert "Rows kept where Condition is drug and Amplitude (pA) < -20." in p.caption


def test_comparisons_at_the_right_level():
    frame = nested(effect=3.0, animals=6, cells=4, seed=3)
    r = Recipe(question="compare", y="Amplitude (pA)", group="Condition", cell="Cell", animal="Animal", stats=True)
    p = prepare.prepare(frame, r)
    (c,) = stats.compare(p.points)
    truth = p.points[p.points.group == "drug"].y.mean() - p.points[p.points.group == "control"].y.mean()
    assert c.difference == pytest.approx(truth) and c.ci[0] < truth < c.ci[1]
    assert c.resampled == "animals, then cells within them" and "24 and 24 cells from 6 and 6 animals" in c.n
    assert "Mixed model" in c.mixed and "Hedges' g" in c.sentence("pA")
    assert "hierarchical bootstrap" in stats.methods([c])
    # paired: the same cells before and after
    before = pd.DataFrame({"Cell": [f"c{k}" for k in range(8)], "Phase": "before", "Y": np.arange(8.0)})
    after = before.assign(Phase="after", Y=before.Y + 2.0 + np.linspace(-0.1, 0.1, 8))
    pp = prepare.prepare(pd.concat([before, after]), Recipe(question="compare", y="Y", group="Phase", cell="Cell",
                                                            paired=True))
    (c,) = stats.compare(pp.points, paired=True)
    assert c.paired and c.difference == pytest.approx(2.0) and c.test == "paired t test"
    assert stats.holm([0.01, 0.04, 0.03]) == pytest.approx([0.03, 0.06, 0.06])


def test_pooling_events_inflates_false_positives_but_animals_do_not():
    rate = stats.nested_false_positives(n_sims=300)
    assert 0.0 <= rate["animals"] <= 0.10
    assert rate["pooled events"] > 0.4 > rate["animals"]


@pytest.mark.parametrize("ext", ["pdf", "svg", "eps", "png", "tiff"])
def test_every_export_reopens_editable(tmp_path, ext):
    frame = nested()
    spec = Spec([Recipe(data=pack(frame), question="compare", plot="bar", y="Amplitude (pA)", group="Condition",
                        cell="Cell", animal="Animal", stats=True),
                 Recipe(data=pack(frame), question="distribution", plot="cumulative", y="Amplitude (pA)",
                        group="Condition", cell="Cell", level="row")], columns=2, size="double")
    path = render.export(spec, tmp_path / f"fig.{ext}", dpi=150)
    back = render.read_spec(path)
    assert back.to_json() == spec.to_json()
    assert back.panels[0].frame().shape == frame.shape


def test_every_plot_type_draws():
    frame = nested()
    t = pd.DataFrame([{"Cell": c, "Condition": g, "Time (min)": x, "Slope": 1 + 0.01 * x + c}
                      for g in ("a", "b") for c in range(3) for x in range(30)])
    cases = [(frame, dict(question="compare", plot=p, y="Amplitude (pA)", group="Condition", cell="Cell"))
             for p in ("points", "bar", "box", "violin", "estimation", "superplot")]
    cases += [(t, dict(question="time", plot=p, x="Time (min)", y="Slope", group="Condition", cell="Cell",
                       normalize=[0, 5], bin=3)) for p in ("time_course", "lines")]
    cases += [(t, dict(question="relationship", plot=p, x="Time (min)", y="Slope", group="Condition", fit=f,
                       level="row", y2="Slope" if p == "line" else ""))
              for p, f in (("scatter", "linear"), ("line", ""), ("xy_error", ""))]
    cases += [(t.assign(Z=t.Slope * 2), dict(question="relationship", plot="heatmap", x="Time (min)", y="Cell",
                                             y2="Z", level="row"))]
    cases += [(frame, dict(question="distribution", plot=p, y="Amplitude (pA)", group="Condition", cell="Cell",
                           level="row")) for p in ("cumulative", "histogram")]
    trace = pd.DataFrame({"Time (ms)": np.arange(500) * 0.1, "Current (pA)": -50 * np.exp(-np.arange(500) / 80.0)})
    cases += [(trace, dict(question="example", plot="trace", x="Time (ms)", y="Current (pA)", level="row",
                           scale_bars=True))]
    for frame_, kw in cases:
        fig, drawn = render.build(Spec([Recipe(data=pack(frame_), **kw)]))
        assert drawn[0].prepared is not None and fig.axes, kw


def test_tables_from_excel_and_from_a_clampwise_workbook(tmp_path):
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "My data"
    ws.append(["Cell", "Condition", "Peak (pA)"])
    for k in range(6):
        ws.append([f"c{k}", "control" if k < 3 else "drug", -100 - 10 * k])
    wb.save(tmp_path / "mine.xlsx")
    assert data.sheets(tmp_path / "mine.xlsx") == ["My data"]
    f = data.read_table(tmp_path / "mine.xlsx")
    assert list(f.columns) == ["Cell", "Condition", "Peak (pA)"] and f["Peak (pA)"].dtype.kind in "if"
    assert data.guess(f)["group"] == "Condition" and data.guess(f)["cell"] == "Cell"
    wb = Workbook()
    ws = wb.active
    ws.title = "DSI"
    ws.append(["Recording", "", "DSI"])
    ws.append(["File", "Cell", "Suppression"])
    ws.append(["", "", "%"])
    ws.append(["a.abf", "c1", 45.0])
    wb.create_sheet("Definitions")
    wb.save(tmp_path / "cw.xlsx")
    f = data.read_table(tmp_path / "cw.xlsx", "DSI")
    assert list(f.columns) == ["File", "Cell", "Suppression (%)"] and f.iloc[0, 2] == 45.0
    wide = pd.DataFrame({"Cell": ["a", "b"], "control": [1.0, 2.0], "drug": [3.0, 4.0]})
    long = data.to_long(wide, ["control", "drug"])
    assert len(long) == 4 and set(long.Condition) == {"control", "drug"}
    assert data.to_wide(long, ["Cell"], "Condition", "Value").shape == (2, 3)


def test_a_digitized_plot_recovers_a_known_curve_within_2_percent(tmp_path):
    from matplotlib.figure import Figure

    from clampwise.figure import digitize

    x = np.linspace(0, 10, 400)
    truth = lambda v: 10 * np.exp(-v / 3) + 1                                  # noqa: E731
    fig = Figure(figsize=(5, 4), dpi=100)
    ax = fig.add_subplot()
    ax.plot(x, truth(x), color="#d62728", lw=1.5)
    ax.set_xlim(0, 10)
    ax.set_ylim(0, 12)
    fig.savefig(tmp_path / "plot.png", dpi=100)
    img = digitize.load(tmp_path / "plot.png")
    h = img.shape[0]
    # the axes' known values where a user would click them: 0 and 10 on x, 0 and 12 on y
    (px0, py0), (px1, py1) = ax.transData.transform([(0, 0), (10, 12)])
    cal = digitize.Calibration(digitize.Axis(px0, 0, px1, 10), digitize.Axis(h - py0, 0, h - py1, 12))
    cols, rows = digitize.trace(img, (214, 39, 40), tolerance=60)
    t = digitize.table(cal, cols, rows, "Time (s)", "Value")
    inside = (t["Time (s)"] > 0.2) & (t["Time (s)"] < 9.8)
    err = np.abs(t.loc[inside, "Value"] - truth(t.loc[inside, "Time (s)"])) / (12 - 0)
    assert inside.sum() > 300 and err.max() < 0.02
    assert "estimate" in t.attrs["source"]
    log = digitize.Axis(0, 1, 100, 1000, log=True)
    assert log.value(50) == pytest.approx(10 ** 1.5)


def test_the_dsi_paper_figure_and_its_methods():
    from clampwise import simulate
    from clampwise.analysis import pipeline
    from clampwise.figure import layouts, methods
    from clampwise.project.presets import BUILTIN

    sim = simulate.dsi_evoked(durations=(1.0, 5.0), trials=2)
    run = pipeline.analyze(sim.recording, pipeline.Settings(by_analysis=BUILTIN["Evoked IPSC DSI"].settings))
    spec = layouts.dsi_figure(run.results["dsi"], lambda name: sim.recording)
    assert spec is not None and len(spec.panels) == 3 and spec.size == "double"
    a = spec.panels[0].frame()
    assert set(a.Trace) == {"before the step", "just after the step"} and spec.panels[0].source["kind"] == "traces"
    fig, drawn = render.build(spec)
    assert all(d.prepared is not None for d in drawn)
    assert "episode" in drawn[1].prepared.caption and "cell" not in drawn[1].prepared.caption
    text = methods.paragraph(drawn)
    assert "(A)" in text and "(C)" in text
    traces = [np.zeros(10), np.ones(10), np.full(10, 0.4), np.full(10, 5.0)]
    assert layouts.closest_to_mean(traces) == 1                                  # mean 1.6: 1.0 is nearest


def test_a_real_world_sheet_two_tables_side_by_side_with_their_summaries(tmp_path):
    """Names in row 4 under a title, two groups side by side with the same headers,
    AVERAGE/SEM/n rows under them, and cell names that repeat across animals."""
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "decay"
    ws.append([None, "Decay times"])
    ws.append([])
    ws.append([None, "resistant", None, None, None, "sensitive"])
    ws.append(["Subject", "Slice#,cell#", "Decay (ms)", None, "Subject", "Slice#,cell#", "Decay (ms)"])
    rows = [("A1", "1,1", 50.0, "B1", "1,1", 80.0), ("A1", "1,2", 60.0, "B2", "1,1", 90.0),
            ("A2", "1,1", 70.0, "B3", "2,1", 100.0), (None, "AVERAGE", 60.0, None, "AVERAGE", 90.0),
            (None, "SEM", 5.8, None, "SEM", 5.8)]
    for a, b, c, d, e, f in rows:
        ws.append([a, b, c, None, d, e, f])
    wb.save(tmp_path / "sheet.xlsx")
    assert data.sniff(tmp_path / "sheet.xlsx", "decay") == (3, 0)
    f = data.read_table(tmp_path / "sheet.xlsx", "decay")
    long = data.to_long(f, ["Decay (ms)", "Decay (ms) (2)"], "Cells", "Decay (ms)", ["resistant", "sensitive"])
    assert list(long[long.Cells == "sensitive"].Subject.dropna()) == ["B1", "B2", "B3"]    # its own block's subjects
    g = data.guess(long)
    assert (g["cell"], g["animal"]) == ("Slice#,cell#", "Subject")
    r = Recipe(question="compare", y="Decay (ms)", group="Cells", cell=g["cell"], animal=g["animal"])
    p = prepare.prepare(long, r)
    assert any("summaries" in w for w in p.warnings)
    r.filters = [["Subject", "has a value", ""]]
    p = prepare.prepare(long, r)
    assert not p.warnings and p.counts["resistant"][:2] == (3, 2) and p.counts["sensitive"][:2] == (3, 3)
    # "1,1" in A1 and in A2 are two cells: a cell's name is unique only within its animal
    assert sorted(p.points[p.points.group == "resistant"].y) == [50.0, 60.0, 70.0]


def test_a_time_course_says_when_cells_do_not_line_up():
    rows = [{"Cell": c, "t (s)": k * 4.0 + 0.1 * c, "Y": 1.0} for c in range(4) for k in range(10)]
    r = Recipe(question="time", plot="time_course", x="t (s)", y="Y", cell="Cell")
    assert any("do not line up" in w and "about 4 s" in w for w in prepare.prepare(pd.DataFrame(rows), r).warnings)
    r.bin = 4.0
    assert not prepare.prepare(pd.DataFrame(rows), r).warnings


def test_themes_group_looks_order_and_hidden_groups():
    from clampwise.figure.recipe import THEMES, Style

    st = Style()
    st.apply_theme("Dark slide")
    assert (st.background, st.ink, st.theme) == ("#1E1E1E", "#E8E8E8", "Dark slide") and set(THEMES) >= {
        "Journal (classic)", "Greyscale print"}
    params = render.rc(st)
    assert params["axes.facecolor"] == "#1E1E1E" and params["text.color"] == "#E8E8E8"
    frame = nested()
    r = Recipe(question="compare", y="Amplitude (pA)", group="Condition", cell="Cell", order=["drug", "control"],
               series={"control": {"colour": "#123456", "marker": "s", "filled": False, "label": "Ctrl"}})
    lk = render.looks(r, st, render.ordered(["control", "drug"], r))
    assert list(lk) == ["drug", "control"] and lk["control"].colour == "#123456" and not lk["control"].filled
    assert lk["control"].label == "Ctrl" and lk["drug"].marker == "o"
    hidden = Recipe(**{**r.__dict__, "series": {"drug": {"hidden": True}}})
    p = prepare.prepare(frame, hidden)
    assert set(p.points.group) == {"control"} and "Not shown: drug." in p.caption
    fig, drawn = render.build(Spec([Recipe(**{**r.__dict__, "data": pack(frame)})], style=st))
    assert [t.get_text() for t in fig.axes[0].get_xticklabels()] == ["drug", "Ctrl"]


def test_counts_and_percentages():
    cells = pd.DataFrame({"Cell": [f"c{k}" for k in range(10)] * 2, "Group": ["a"] * 10 + ["b"] * 10,
                          "Fires": ["yes"] * 3 + ["no"] * 7 + ["yes"] * 6 + ["no"] * 4})
    r = Recipe(question="counts", plot="count_bar", x="Fires", group="Group", cell="Cell")
    p = prepare.prepare(cells.assign(Cell=cells.Group + cells.Cell), r)
    got = {(g, c): n for g, c, n in p.points[["group", "xcat", "count"]].itertuples(index=False)}
    assert got == {("a", "yes"): 3, ("a", "no"): 7, ("b", "yes"): 6, ("b", "no"): 4}
    r.percent = True
    p = prepare.prepare(cells.assign(Cell=cells.Group + cells.Cell), r)
    assert p.points.set_index(["group", "xcat"]).y[("b", "yes")] == pytest.approx(60.0)
    for plot in ("count_bar", "stacked_bar", "pie", "donut"):
        fig, drawn = render.build(Spec([Recipe(**{**r.__dict__, "plot": plot, "data": pack(cells)})]))
        assert drawn[0].prepared is not None, plot


def test_a_regression_band_and_every_new_plot_type_draws():
    x = np.linspace(0, 10, 50)
    grid, fit, lo, hi, eq = render.regression_band(x, 2 * x + 1)
    assert fit[0] == pytest.approx(1.0) and fit[-1] == pytest.approx(21.0) and np.allclose(hi - lo, 0, atol=1e-6)
    noisy = 2 * x + 1 + np.random.default_rng(0).normal(0, 1, x.size)
    _, fit, lo, hi, _ = render.regression_band(x, noisy)
    assert (lo < fit).all() and (hi > fit).all() and (hi - lo)[25] < (hi - lo)[0]     # narrowest at the middle
    frame = nested()
    frame["Sex"] = np.where(frame.Animal.str[-1].astype(int) % 2 == 0, "f", "m")
    xy = pd.DataFrame({"X": x, "Y": noisy, "S": np.abs(noisy)})
    cases = [(frame, dict(question="compare", plot=p, y="Amplitude (pA)", group="Condition", cell="Cell",
                          x="Sex")) for p in ("raincloud", "dot", "hbar", "grouped_bar")]
    cases += [(xy, dict(question="relationship", plot=p, x="X", y="Y", size_by="S", level="row",
                        ref_lines=[{"axis": "y", "at": 5, "label": "five"}]))
              for p in ("regression", "step", "bubble", "hexbin")]
    cases += [(frame, dict(question="distribution", plot=p, y="Amplitude (pA)", group="Condition", level="row"))
              for p in ("density", "qq")]
    t = pd.DataFrame([{"Cell": c, "T": k, "Y": 1.0 + k} for c in range(3) for k in range(10)])
    cases += [(t, dict(question="time", plot="area", x="T", y="Y", cell="Cell", error_style="band"))]
    for f, kw in cases:
        fig, drawn = render.build(Spec([Recipe(data=pack(f), **kw)]))
        assert drawn[0].prepared is not None and fig.axes, kw["plot"]


def test_trace_sets_averaged_each_or_typical_and_dsi_roles():
    """Raw traces picked by hand: a set drawn as its average, every sweep, or the sweep closest to their mean;
    DSI responses picked by their place around the steps (an episode's baseline, the first after its step)."""
    from clampwise import simulate
    from clampwise.analysis import pipeline
    from clampwise.figure import traces
    from clampwise.project.presets import BUILTIN

    sim = simulate.dsi_evoked(durations=(1.0, 5.0), trials=2)
    run = pipeline.analyze(sim.recording, pipeline.Settings(by_analysis=BUILTIN["Evoked IPSC DSI"].settings))
    roles = traces.dsi_roles(run.results["dsi"])
    n_ep = len(run.results["dsi"].tables["episodes"])
    pre, post = traces.dsi_pick(roles, 0, before=4), traces.dsi_pick(roles, 0, after=3)
    assert len(pre) == 4 * n_ep and len(post) == 3 * n_ep and not {tuple(i) for i in pre} & {tuple(i) for i in post}
    assert all("before the step" in roles[i]["label"] for i in traces.dsi_indices(roles, 1, before=2))
    sets = [traces.TraceSet("Before", "average", pre), traces.TraceSet("After", "each", post),
            traces.TraceSet("Typical", "typical", pre)]
    w = traces.Window(5.0, 40.0, True, traces.dsi_blank(roles))
    frame, notes, markers = traces.build(sets, lambda _f: sim.recording, w)
    assert list(frame.Trace.unique()) == ["Before", "After", "Typical"]
    assert frame[frame.Trace == "After"].Sweep.nunique() == len(post) and frame[frame.Trace == "Before"].Sweep.nunique() == 1
    t = frame["Time (ms)"]
    assert t.min() >= -5.01 and t.max() < 40.0 and frame[frame.Trace == "Before"].iloc[:, 1].abs()[t < -0.5].max() < 1e3
    peak = lambda name: float(frame[(frame.Trace == name) & (t > w.blank_ms)].iloc[:, 1].abs().max())   # noqa: E731
    assert peak("After") < peak("Before")                       # suppressed after the step
    assert any("average of" in n for n in notes) and (markers[0]["end"] == w.blank_ms if w.blank_ms else True)
    src = traces.source(sets, w, "x")
    back, w2 = traces.from_source(src)
    assert [s.name for s in back] == ["Before", "After", "Typical"] and w2 == w
    from clampwise.figure import layouts
    from clampwise.figure.recipe import Spec

    fig, drawn = render.build(Spec([layouts.example_recipe(sets, w, lambda _f: sim.recording, "x")]))
    assert drawn[0].prepared is not None and not drawn[0].prepared.warnings
