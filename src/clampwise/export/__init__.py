"""Exports (B14): a project's numbers out, every column defined.

One workbook per project, built from the registry, nothing written per measure:
- a sheet per recording type: who the recording is of, its type, quality and
  flags, then every recording-level output of the analyses that ran, with its
  units; each column's definition and source is a comment on its header (as
  NeuroExpress explains its parameters), and each analysis's choices (which
  pulse, window or baseline stands for the recording) sit beside its numbers;
- the cells' results (DSI from a cell's events), when there are any;
- every result table (steps, pulses, events, ...) on its own sheet;
- Definitions (every analysis and output) and Provenance (where each number came
  from: file hash, version, settings changed from the defaults).

The same numbers as tidy CSV (one value per row) with the tables, the
definitions and a manifest. Results are read from the project's store; no
recording is re-analyzed or even opened. openpyxl only.
"""

from __future__ import annotations

import csv
import datetime as dt
import json
import math
import re
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Iterable

from .. import __version__
from ..analysis import registry

IDENTITY = ("File", "Animal", "Slice", "Cell", "Condition", "Age", "Sex", "Genotype", "Region", "Cell type", "Type",
            "Quality", "Flags", "Review")
IDENTITY_HELP = {
    "File": "The recording's file.",
    "Animal": "The animal: read from the folder names by the project's naming rules, or typed (Files, Details).",
    "Slice": "The slice, the same way.",
    "Cell": "The cell, the same way.",
    "Condition": "The condition (control, a drug...), the same way.",
    "Age": "Typed on Files (experiment details).", "Sex": "Typed on Files (experiment details).",
    "Genotype": "Typed on Files (experiment details).", "Region": "Typed on Files (experiment details).",
    "Cell type": "Typed on Files (experiment details).",
    "Quality": "passed: no quality flag; flagged: at least one (see Flags); not checked: no quality analysis ran.",
    "Flags": "Every quality flag of the recording: what is wrong, the limit crossed and where the limit comes from. "
             "Flags are evidence; no number is changed or dropped because of one.",
    "Review": "good or bad: marked by hand on Check (G or B), with the reason typed; blank: not reviewed. A mark is "
              "the user's word; no number is changed or dropped because of one.",
    "Type": "The recording type the pipeline decided, which chose the analyses that ran (Check shows why).",
}


@dataclass
class Exported:
    """One recording's stored results, as exported."""

    path: Path
    name: str
    identity: Any
    metadata: Any
    recording_type: str
    results: dict
    provenance: Any
    review: str = ""


def collect(project) -> list[Exported]:
    """Every analyzed recording of ``project``, from its store."""
    out = []
    for row in project.recordings():
        run = project.last_run(row.path, read_file=False)
        if run is not None:
            out.append(Exported(row.path, row.name, project.identity(row.path), project.metadata(row.path),
                                run.recording_type, run.results, run.provenance, _review(project, row.path)))
    return out


def _review(project, path) -> str:
    verdict, reason = project.file_review(path)
    return f"{verdict}: {reason}" if verdict and reason else verdict


def cell_results(project, rows: list[Exported]) -> list:
    """The stored cell-level results, one per cell, in the order of the cells' first files."""
    seen, out = set(), []
    for r in rows:
        c = project.cell_result(r.path)
        if c is not None and c.key not in seen and c.results:
            seen.add(c.key)
            out.append(c)
    return out


# ------------------------------------------------------------------ what goes in the columns


def plain(v: Any) -> Any:
    """A value as a spreadsheet takes it: numbers as numbers, NaN as an empty cell."""
    if v is None:
        return None
    if hasattr(v, "item") and not isinstance(v, (list, tuple, dict)):
        try:
            v = v.item()
        except (ValueError, AttributeError):
            pass
    if isinstance(v, bool):
        return "yes" if v else "no"
    if isinstance(v, float):
        return None if not math.isfinite(v) else v
    if isinstance(v, (int, str)):
        return v
    if isinstance(v, dict):
        return "; ".join(f"{k}: {plain(x)}" for k, x in v.items())
    if isinstance(v, (list, tuple)):
        return ", ".join(str(plain(x)) for x in v)
    return str(v)


def describe(o: registry.Output | None, analysis: str) -> str:
    """The comment on a column's header."""
    a = registry.load().get(analysis)
    where = a.label if a else analysis
    if o is None:
        return f"{where}: no definition is registered for this column yet."
    text = o.definition + (f"\nSource: {o.source}" if o.source else "")
    return f"{text}\n({where}; {o.units or 'no units'})"


@dataclass
class Column:
    group: str                          # the analysis's label over its first column
    title: str
    units: str
    comment: str
    get: Any                            # Exported -> value


def identity_columns() -> list[Column]:
    def ident(attr):
        return lambda r: getattr(r.identity, attr, "")

    def meta(attr):
        return lambda r: getattr(r.metadata, attr, "") if r.metadata is not None else ""

    def quality(r):
        q = r.results.get("quality")
        return "not checked" if q is None else ("flagged" if q.qc_flags else "passed")

    def flags(r):
        out = [f"{registry.load()[k].label if k in registry.load() else k}: software error ({res.error.splitlines()[0]})"
               for k, res in r.results.items() if res.error]
        return "; ".join(out + [f.describe() for res in r.results.values() for f in res.qc_flags])

    get = {"File": lambda r: r.name, "Animal": ident("animal"), "Slice": ident("slice"), "Cell": ident("cell"),
           "Condition": ident("condition"), "Age": meta("age"), "Sex": meta("sex"), "Genotype": meta("genotype"),
           "Region": meta("region"), "Cell type": meta("cell_type"),
           "Type": lambda r: registry.TYPE_WORDS.get(r.recording_type, r.recording_type), "Quality": quality,
           "Flags": flags, "Review": lambda r: r.review}
    return [Column("Recording" if k == 0 else "", t, "", IDENTITY_HELP.get(t, ""), get[t]) for k, t in enumerate(IDENTITY)]


def value_columns(results: Iterable[dict], level: str = "recording") -> list[Column]:
    """Every output of the analyses present in ``results`` (dicts of key -> Result),
    in registry order: the declared outputs at ``level`` first, then any other
    value an analysis gave, then its choices."""
    results = list(results)
    cols = []
    for key, a in registry.load().items():
        present = [r[key] for r in results if key in r]
        if not present:
            continue
        given = {k for res in present for k in res.values}
        declared = [o for o in a.outputs if o.level == level and o.key in given]
        others = sorted(given - {o.key for o in declared} - {o.key for o in a.outputs if o.level != level})
        start = len(cols)
        for o in declared:
            cols.append(Column("", o.title, o.units, describe(o, key),
                               lambda r, k=key, v=o.key: plain(r.results[k].values.get(v)) if k in r.results else None))
        for v in others:
            cols.append(Column("", v.replace("_", " ").capitalize(), "", describe(None, key),
                               lambda r, k=key, v=v: plain(r.results[k].values.get(v)) if k in r.results else None))
        if any(res.choices for res in present):
            cols.append(Column("", "Choices", "", f"{a.label}: which pulse, sweep, window or baseline stands for the "
                               "recording, and why (nothing is chosen silently).",
                               lambda r, k=key: plain(r.results[k].choices) if k in r.results else None))
        if len(cols) > start:
            cols[start].group = a.label
    return cols


def tables(rows: list[Exported], cells: Iterable = ()) -> list[tuple[str, str, list[str], list[tuple[Any, dict]]]]:
    """Every result table across the recordings, then the cells' (their File is
    the cell): (analysis, table, columns, [(recording or cell, row)])."""
    from types import SimpleNamespace

    found: dict[tuple[str, str], tuple[list[str], list]] = {}
    sources = list(rows) + [SimpleNamespace(name=c.key, results=c.results) for c in cells]
    for key in registry.load():
        for r in sources:
            res = r.results.get(key)
            if res is None:
                continue
            for name, table in res.tables.items():
                cols, body = found.setdefault((key, name), ([], []))
                for row in table:
                    for c in row:
                        if c not in cols:
                            cols.append(c)
                    body.append((r, row))
    return [(k, n, cols, body) for (k, n), (cols, body) in found.items() if body]


def definitions() -> list[dict]:
    from ..analysis import columns

    out = []
    for key, a in registry.load().items():
        out.append({"Analysis": a.label, "Key": key, "Name": a.label, "Units": "", "Level": f"analysis, v{a.version}",
                    "Definition": a.description, "Reference": a.source, "Verified against": ""})
        for o in a.outputs:
            out.append({"Analysis": a.label, "Key": f"{key}.{o.key}", "Name": o.title, "Units": o.units, "Level": o.level,
                        "Definition": o.definition, "Reference": o.source, "Verified against": o.verified_against})
        for o in columns.TABLES.get(key, ()):
            out.append({"Analysis": a.label, "Key": f"{key}.{o.key}", "Name": o.title, "Units": o.units,
                        "Level": "table column", "Definition": o.definition, "Reference": o.source,
                        "Verified against": ""})
    for o in columns.COMMON.values():
        out.append({"Analysis": "(every table)", "Key": o.key, "Name": o.title, "Units": o.units,
                    "Level": "table column", "Definition": o.definition, "Reference": o.source, "Verified against": ""})
    return out


def provenance(rows: list[Exported], features: str = "") -> list[dict]:
    out = []
    for r in rows:
        p = r.provenance
        changed = [f"{k}.{s}={plain(res.settings.get(s))}" for k, res in r.results.items() for s in res.changed_settings]
        out.append({"File": r.name, "Path": str(r.path), "SHA-256": getattr(p, "content_hash", ""),
                    "Clampwise": getattr(p, "app_version", ""), "Analyzed at": getattr(p, "analyzed_at", ""),
                    "Preset": getattr(p, "preset", ""),
                    "Analyses (version)": plain(getattr(p, "analyses", "")),
                    "Settings changed from the defaults": "; ".join(changed), "Features": features})
    return out


def _features(project) -> str:
    from ..project.features import describe as describe_features

    return describe_features(project.features())


# ------------------------------------------------------------------ the workbook


def _sheet_name(text: str, taken: set[str]) -> str:
    base = re.sub(r"[\[\]:*?/\\]", " ", text).strip()[:31] or "Sheet"
    name, k = base, 2
    while name.lower() in taken:
        name = f"{base[:26]} ({k})"
        k += 1
    taken.add(name.lower())
    return name


EXCEL_ROWS = 1_048_576 - 3        # a sheet's rows, less the three header rows


def write_workbook(project, path: str | Path | None = None) -> Path:
    """The project's workbook (default: in its folder, named for it and today);
    a file Excel holds open is left alone and the next free name taken."""
    from openpyxl import Workbook
    from openpyxl.cell import WriteOnlyCell
    from openpyxl.comments import Comment
    from openpyxl.styles import Alignment, Font
    from openpyxl.utils import get_column_letter

    rows = collect(project)
    path = Path(path) if path else Path(project.folder) / f"{project.name} - Clampwise {dt.date.today():%Y-%m-%d}.xlsx"
    book = Workbook(write_only=True)
    taken: set[str] = set()
    bold, group_font, wrap = Font(bold=True), Font(bold=True, italic=True), Alignment(wrap_text=True, vertical="top")

    def cell(ws, value, font=None, comment=""):
        c = WriteOnlyCell(ws, value=value)
        if isinstance(value, str) and value.startswith("="):
            c.data_type = "s"                   # text, never a formula
        if font is not None:
            c.font = font
        if comment:
            c.comment = Comment(comment, "Clampwise", width=320, height=min(400, 40 + 15 * (len(comment) // 45)))
        return c

    def sheet(title: str, cols: list[Column], items: list, getter=lambda col, item: col.get(item)) -> None:
        if len(items) > EXCEL_ROWS:                 # Excel's limit: the rest on sheets of their own
            for k in range(0, len(items), EXCEL_ROWS):
                sheet(f"{title} {k // EXCEL_ROWS + 1}", cols, items[k:k + EXCEL_ROWS], getter)
            return
        ws = book.create_sheet(_sheet_name(title, taken))
        ws.freeze_panes = "B4"
        for j, c in enumerate(cols, start=1):
            ws.column_dimensions[get_column_letter(j)].width = max(10, min(40, len(c.title) + 4))
        ws.append([cell(ws, c.group, group_font) for c in cols])
        ws.append([cell(ws, c.title, bold, c.comment) for c in cols])
        ws.append([cell(ws, c.units) for c in cols])
        for item in items:
            ws.append([cell(ws, getter(c, item)) for c in cols])

    by_type: dict[str, list[Exported]] = {}
    for r in rows:
        by_type.setdefault(r.recording_type, []).append(r)
    for t in registry.TYPES:
        group = by_type.get(t)
        if group:
            sheet(registry.TYPE_WORDS.get(t, t), identity_columns() + value_columns(r.results for r in group), group)
    cells = cell_results(project, rows)
    if cells:
        cols = [Column("Cell", "Cell", "", "The cell: animal, slice and cell.", lambda c: c.key),
                Column("", "Files", "", "The cell's recordings these results combine.", lambda c: ", ".join(c.files))]
        sheet("Cells", cols + value_columns((c.results for c in cells), level="cell"), cells)
    for key, name, cols, body in tables(rows, cells):
        a = registry.load()[key]
        outs = {c: registry.column(key, c) for c in cols}
        tcols = [Column(a.label, "File", "", "The recording's file.", lambda item: item[0].name)] + [
            Column("", outs[c].title if outs[c] else c.replace("_", " ").capitalize(), outs[c].units if outs[c] else "",
                   describe(outs[c], key), lambda item, c=c: plain(item[1].get(c))) for c in cols]
        sheet(name if name == key else f"{name} ({key})", tcols, body)
    for title, table in (("Definitions", definitions()), ("Provenance", provenance(rows, _features(project)))):
        ws = book.create_sheet(_sheet_name(title, taken))
        ws.freeze_panes = "A2"
        heads = list(table[0]) if table else ["(nothing analyzed)"]
        for j, h in enumerate(heads, start=1):
            ws.column_dimensions[get_column_letter(j)].width = 70 if h == "Definition" else 22
        ws.append([cell(ws, h, bold) for h in heads])
        for row in table:
            out = [cell(ws, plain(row[h])) for h in heads]
            for c in out:
                c.alignment = wrap
            ws.append(out)
    path.parent.mkdir(parents=True, exist_ok=True)
    target, k = path, 2
    while True:
        try:
            book.save(target)
            project.set_meta("last_export", dt.datetime.now().isoformat(timespec="seconds"))
            return target
        except PermissionError:                 # open in Excel: keep it, write beside it
            target = path.with_name(f"{path.stem} ({k}){path.suffix}")
            k += 1


# ------------------------------------------------------------------ tidy CSV


def write_csv(project, folder: str | Path | None = None) -> list[Path]:
    """Tidy CSV: every value one row (file, who, type, analysis, output, value,
    units), each result table its own file, the definitions, and a manifest."""
    rows = collect(project)
    folder = Path(folder) if folder else Path(project.folder) / f"{project.name} - Clampwise CSV"
    folder.mkdir(parents=True, exist_ok=True)
    written = []

    def save(name: str, heads: list[str], body: Iterable[list]) -> None:
        p = folder / name
        with open(p, "w", newline="", encoding="utf-8-sig") as f:
            w = csv.writer(f)
            w.writerow(heads)
            w.writerows(body)
        written.append(p)

    ident = identity_columns()
    heads = [c.title for c in ident if c.title != "Flags"] + ["Analysis", "Output", "Value", "Units"]
    body = []
    for r in rows:
        who = [c.get(r) for c in ident if c.title != "Flags"]
        for key, res in r.results.items():
            a = registry.load().get(key)
            for v, x in res.values.items():
                o = a.output(v) if a else None
                body.append(who + [key, v, plain(x), o.units if o else ""])
    save("values.csv", heads, body)
    save("flags.csv", ["File", "Analysis", "Flag", "What", "Limit", "Reference", "Where"],
         [[r.name, k, f.code, f.text, f.limit, f.source, f.where] for r in rows for k, res in r.results.items()
          for f in res.qc_flags])
    for key, name, cols, tbody in tables(rows, cell_results(project, rows)):
        save(f"table {key} {name}.csv", ["File"] + cols, [[r.name] + [plain(row.get(c)) for c in cols] for r, row in tbody])
    defs = definitions()
    save("definitions.csv", list(defs[0]), [list(d.values()) for d in defs])
    prov = provenance(rows, _features(project))
    if prov:
        save("provenance.csv", list(prov[0]), [list(p.values()) for p in prov])
    manifest = folder / "manifest.json"
    manifest.write_text(json.dumps({
        "written_by": f"Clampwise {__version__}", "written_at": dt.datetime.now().isoformat(timespec="seconds"),
        "project": str(project.folder), "recordings": len(rows),
        "files": {p.name: ("one row per value: who, type, analysis, output, value, units" if p.name == "values.csv"
                           else "every column's definition, units and source" if p.name == "definitions.csv"
                           else "every quality flag" if p.name == "flags.csv"
                           else "where each number came from" if p.name == "provenance.csv"
                           else "a result table, one row per item") for p in written}}, indent=1), encoding="utf-8")
    project.set_meta("last_export", dt.datetime.now().isoformat(timespec="seconds"))
    return written + [manifest]
